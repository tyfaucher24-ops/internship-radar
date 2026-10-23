"""Builds the V14.3 synthetic resume fixtures: tests/resume-fixtures/<id>/{resume.pdf,resume.docx,[resume.txt],expected.json}.
Usage: python3 tests/resume-fixtures/build.py      (needs reportlab, python-docx; LibreOffice `soffice` for the Word-format PDFs)
All people are fictional (see people.py). expected.json is written from the spec below, never from parser output."""
def ityp(t):
    import re as _r
    return 'Co-op' if _r.search(r'co-?op',t,_r.I) else 'Internship' if _r.search(r'intern|extern|summer (analyst|associate|scholar|fellow)',t,_r.I) else 'Job'
import os,sys,json,random,shutil,subprocess,copy
sys.path.insert(0,os.path.dirname(__file__))
from people import *
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor,white,black
import docx
from docx.shared import Pt,Inches,RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
OUT=os.path.dirname(os.path.abspath(__file__))
FD='/usr/share/fonts/truetype/'
for n,f in [('LS','liberation/LiberationSans-Regular.ttf'),('LSB','liberation/LiberationSans-Bold.ttf'),('LSI','liberation/LiberationSans-Italic.ttf'),('SR','dejavu/DejaVuSerif.ttf'),('SRB','dejavu/DejaVuSerif-Bold.ttf'),('SRI','liberation/LiberationSerif-Italic.ttf'),('MO','dejavu/DejaVuSansMono.ttf'),('MOB','dejavu/DejaVuSansMono-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(n,FD+f))
MS=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'];ML=['January','February','March','April','May','June','July','August','September','October','November','December']
TY,TM,_=TODAY
def ym(y,m):return f'{y}-{m:02d}'
def months(sy,sm,ey,em):return (ey-sy)*12+(em-sm)+1
def fmt_dates(d,style):
    """d = dict(sy,sm,ey,em,cur,season). Returns (printed, truth{start,end,months})"""
    sy,sm,ey,em,cur=d['sy'],d['sm'],d.get('ey'),d.get('em'),d.get('cur',False)
    end_m=lambda: (TY,TM) if cur else (ey,em)
    if d.get('season') and style in('season','mon','full'):
        s=d['season'];p=f'{s} {sy}';b={'Summer':(6,8),'Fall':(9,12),'Spring':(1,5)}[s]
        return p,dict(start=ym(sy,b[0]),end=ym(sy,b[1]),months=b[1]-b[0]+1)
    if style=='year':
        p=f'{sy} – '+('Present' if cur else str(ey));return p,dict(start=str(sy),end='Present' if cur else str(ey),months=None)
    if style=='num':p=f'{sm:02d}/{sy} – '+('Present' if cur else f'{em:02d}/{ey}')
    elif style=='full':p=f'{ML[sm-1]} {sy} – '+('Present' if cur else f'{ML[em-1]} {ey}')
    elif style=='to':p=f'{ML[sm-1]} {sy} to '+('Present' if cur else f'{ML[em-1]} {ey}')
    elif style=='apos':p=f"{MS[sm-1]} '{str(sy)[2:]} – "+('Present' if cur else f"{MS[em-1]} '{str(ey)[2:]}")
    elif style=='monrange' and not cur and sy==ey:p=f'{MS[sm-1]}.–{MS[em-1]}. {sy}'
    elif style=='sept':p=('Sept.' if sm==9 else MS[sm-1]+'.')+f' {sy} – '+('Current' if cur else ('Sept.' if em==9 else MS[em-1]+'.')+f' {ey}')
    else:p=f'{MS[sm-1]} {sy} – '+('Present' if cur else f'{MS[em-1]} {ey}')
    e=end_m();return p,dict(start=ym(sy,sm),end='Present' if cur else ym(ey,em),months=months(sy,sm,e[0],e[1]))
def classyear(gy,gm):
    ay=TY+(1 if TM>=7 else 0);g=gy+(1 if (gm or 5)>=9 else 0);d=g-ay
    return 'Senior' if d==0 else 'Junior' if d==1 else 'Sophomore' if d==2 else 'Freshman' if d==3 else 'Incoming' if d>3 else 'Graduated'
HEAD={'edu':['EDUCATION','Education','ACADEMIC BACKGROUND','EDUCATION'],'exp':['EXPERIENCE','PROFESSIONAL EXPERIENCE','Work Experience','RELEVANT EXPERIENCE','Experience'],
 'lead':['LEADERSHIP & ACTIVITIES','Leadership','CAMPUS INVOLVEMENT','ACTIVITIES','Leadership Experience'],'ath':['ATHLETICS','Athletics'],'vol':['VOLUNTEER EXPERIENCE','Community Service','VOLUNTEER'],
 'proj':['PROJECTS','Technical Projects','PROJECTS'],'skills':['SKILLS','TECHNICAL SKILLS','Skills & Interests','SKILLS & INTERESTS','Skills'],'cert':['CERTIFICATIONS','Certifications','LICENSES & CERTIFICATIONS'],'pub':['PUBLICATIONS'],'lang':['LANGUAGES']}
def person(i,fmt,field,level,edges):
    r=random.Random(1000+i);P=F[field]
    first,last=FIRST[i%len(FIRST)],LAST[(i*7)%len(LAST)]
    sname=r.choice(list(SCHOOLS));sp,(sc,sst)=SCHOOLS[sname]
    p=dict(id=f'f{i+1:02d}',fmt=fmt,field=field,level=level,edges=edges,name=f'{first} {last}',email=f'{first.lower()}.{last.lower().replace("-","")}@example.com',
           phone=f'({r.choice(["212","312","404","617","704","919","512","206"])}) 555-01{r.randint(10,99)}',city=sc,st=sst,linkedin=f'linkedin.com/in/{first.lower()}{last.lower().replace("-","")}',edu=[],exp=[],skills=[],certs=[],langs=[],interests=[],workauth=('',''),pubs=[],soft=[])
    gy={'freshman':2030,'dense':2027,'career':2027}.get(level,r.choice([2027,2027,2028,2028,2029]));gm=r.choice([5,5,5,12])
    if level=='freshman':gm=5
    # education
    deg_print,deg=r.choice(P['deg']);maj=r.choice(P['majors'])
    if deg=='BSN':maj='Nursing'
    majors=[maj]
    if 'double' in edges:m2=r.choice([m for m in P['majors'] if m!=maj]);majors=[maj,m2]
    minors=[r.choice(['Mathematics','Spanish','Statistics','Economics','Psychology','History'])] if r.random()<0.45 else []
    if minors and minors[0] in majors:minors=[]
    gpa=None if 'nogpa' in edges else f'{r.uniform(3.1,3.98):.2f}' if level!='freshman' else f'{r.uniform(3.4,3.99):.2f}'
    scale='4.0' if gpa and r.random()<0.6 else ''
    if 'scale10' in edges:gpa,scale=f'{r.uniform(8.1,9.6):.1f}','10.0'
    mgpa=f'{r.uniform(3.5,3.99):.2f}' if 'majorgpa' in edges else ''
    sprint=r.choice(sp);school_truth=sname
    if 'ambiguous' in edges:sprint=r.choice(['USC','Queens']);school_truth='';sc,sst={'USC':('Los Angeles','CA'),'Queens':('New York','NY')}[sprint]
    gstyle=r.choice(['expected','plain','classof','range','label']) if 'weirddates' in edges else r.choice(['expected','plain','expected'])
    if gstyle=='classof':gp=f'Class of {gy}';gmt=''
    elif gstyle=='range':gp=f'Aug {gy-4} – {MS[gm-1]} {gy}';gmt=str(gm)
    elif gstyle=='label':gp=f'Expected Graduation: {gm:02d}/{gy}';gmt=str(gm)
    elif gstyle=='plain':gp=f'{ML[gm-1]} {gy}';gmt=str(gm)
    else:gp=f'Expected {MS[gm-1]} {gy}';gmt=str(gm)
    honors=["Dean's List"] if r.random()<0.4 and level!='freshman' else []
    course=r.sample(['Corporate Finance','Data Structures','Microeconomics','Statistics','Organic Chemistry','Marketing Research','Thermodynamics','Public Policy Analysis','Financial Accounting','Algorithms','Anatomy & Physiology','Linear Algebra'],3) if r.random()<0.5 else []
    abroad=('London School of Economics','Study Abroad, Spring 2026') if 'abroad' in edges else None
    e=dict(school_print=sprint,school=school_truth,ambiguous='ambiguous' in edges,loc=f'{sc}, {sst}',deg_print=deg_print,degree=deg,majors=majors,minors=minors,gpa=gpa,scale=scale,mgpa=mgpa,gp=gp,gy=gy,gm=gmt,
           classYear=classyear(gy,int(gmt) if gmt else None),honors=honors,course=course,abroad=abroad,level='bach')
    p['edu'].append(e)
    if level=='transfer':
        ccn=r.choice(list(CC));c2,s2=CC[ccn]
        p['edu'].append(dict(school_print=ccn,school=ccn,ambiguous=False,loc=f'{c2}, {s2}',deg_print='Associate of Arts',degree='AA',majors=['General Studies'],minors=[],gpa=f'{r.uniform(3.5,3.99):.2f}',scale='4.0',mgpa='',gp=f'May {gy-2}',gy=gy-2,gm='5',classYear='Graduated',honors=[],course=[],abroad=None,level='assoc'))
    if level=='freshman':
        hs=r.choice(HS);p['edu'].append(dict(school_print=hs,school=hs,ambiguous=False,loc=f'{r.choice(["Raleigh","Columbus","Austin","Denver"])}, {r.choice(["NC","OH","TX","CO"])}',deg_print='High School Diploma',degree='High School Diploma',majors=[],minors=[],gpa=f'{r.uniform(3.6,4.0):.2f}',scale='4.0',mgpa='',gp=f'June {gy-4}',gy=gy-4,gm='6',classYear='',honors=[],course=[],abroad=None,level='hs'))
    # experience
    dstyle=lambda: r.choice(['num','monrange','year','to','apos','sept','season']) if 'weirddates' in edges else r.choice(['mon','mon','full'])
    def add(sec,typ,org,title,loc,d,bul):p['exp'].append(dict(sec=sec,type=typ,org=org,title=title,loc=loc,d=d,style=dstyle(),bullets=bul))
    yrs_summer=[gy-1,gy-2] if level!='freshman' else []
    n_int={'dense':3,'freshman':0,'career':1}.get(level,r.choice([1,2,2]))
    orgs=r.sample(P['orgs'],min(len(P['orgs']),4))
    for k in range(n_int):
        o,c,s=orgs[k];t=r.choice(P['titles']);y=yrs_summer[k%2]-(k//2)
        if y>TY or (y==TY and TM<6):y=TY-1 if y>=TY else y
        typ=ityp(t)
        d=dict(sy=y,sm=6,ey=y,em=8) if typ!='Co-op' else dict(sy=y,sm=1,ey=y,em=6)
        if typ=='Internship' and r.random()<0.3:d['season']='Summer'
        add('exp',typ,o,t,f'{c}, {s}',d,r.sample(P['bullets'],r.choice([2,3])))
    if level=='career':
        add('exp','Job','FedEx Ground','Operations Coordinator','Memphis, TN',dict(sy=2018,sm=3,ey=2023,em=7),['Scheduled 40 drivers across 2 shifts and cut overtime 15%','Trained 12 new hires on safety and scanning procedures'])
        add('exp','Military / ROTC','U.S. Army','Logistics Specialist','Fort Liberty, NC',dict(sy=2014,sm=6,ey=2018,em=2),['Managed inventory worth $3M for a 120-soldier company'])
    if level=='freshman':
        so,st_=r.choice(SERVICE);add('exp','Service job',so,st_,f'{r.choice(["Raleigh, NC","Columbus, OH","Austin, TX"])}',dict(sy=gy-5,sm=6,ey=gy-4,em=8),['Served 150+ customers per shift during peak hours','Trained 3 new team members'])
        add('lead','Club leadership',f'{e["school_print"].split(",")[0]} {r.choice(["Key Club","National Honor Society"])}' if False else r.choice(['Key Club','National Honor Society']),'President','',dict(sy=gy-6,sm=9,ey=gy-4,em=6),['Organized 6 service events with 80 student volunteers'])
    if level in('standard','dense','transfer','international','athlete') and r.random()<0.6:
        so,st_=r.choice(SERVICE);y=gy-3;add('exp','Service job',so,st_,f'{sc}, {sst}',dict(sy=y,sm=5,ey=y,em=8),['Handled cash and card transactions for 100+ customers daily'])
    if level=='dense' or r.random()<0.35:
        add('exp','Research assistant',f'{e["school_print"].split(",")[0].replace("The ","")} Department of {r.choice(["Economics","Psychology","Biology","Computer Science"])}','Undergraduate Research Assistant',f'{sc}, {sst}',dict(sy=TY-1,sm=9,cur=True),['Cleaned and merged 3 survey datasets with 10,000 rows in Stata','Summarized 25 papers for a literature review'])
    if level=='dense' or r.random()<0.25:
        add('exp','Teaching assistant',f'{e["school_print"].split(",")[0].replace("The ","")} Department of Mathematics','Teaching Assistant',f'{sc}, {sst}',dict(sy=TY,sm=1,cur=True),['Held weekly office hours for 60 students in Calculus I'])
    if r.random()<0.12:add('exp','Founder / own business',f'{last} Tutoring LLC','Founder',f'{sc}, {sst}',dict(sy=gy-3,sm=1,cur=True),['Tutored 25 high school students in SAT math; 90% improved 100+ points'])
    if r.random()<0.12:add('exp','Freelance','Self-Employed','Freelance Graphic Designer','Remote',dict(sy=gy-3,sm=3,ey=gy-2,em=12),['Designed logos and menus for 9 local businesses'])
    if r.random()<0.1:add('exp','Family business',f'{last} Hardware','Operations Assistant',f'{sc}, {sst}',dict(sy=gy-5,sm=6,ey=gy-3,em=8),['Helped run the family business; reorganized inventory of 2,000 SKUs'])
    if r.random()<0.12:add('exp','Part-time job','Campus Recreation Center','Part-Time Front Office Assistant',f'{sc}, {sst}',dict(sy=gy-2,sm=9,ey=gy-1,em=5),['Checked in 300 members per shift and answered phones'])
    # leadership & clubs
    short=e['school_print'].split(',')[0].replace('The ','').replace('University of ','').split(' University')[0]
    if level!='freshman':
        for cl in r.sample(CLUBS,{'dense':3}.get(level,r.choice([1,1,2]))):
            nm,typ,title=cl;add('lead',typ,f'{short} {nm}',title,'',dict(sy=gy-3,sm=9,cur=True) if r.random()<0.6 else dict(sy=gy-3,sm=9,ey=gy-1,em=5),['Recruited 25 new members and ran 10 weekly meetings' if 'leadership' in typ else 'Pitched 2 stocks to a 30-member committee; one was added to the $150K portfolio' if 'fund' in typ else 'Attended weekly workshops on industry topics'])
        if r.random()<0.2:add('lead','Case competition',f'{r.choice(["Deloitte","KPMG","Net Impact"])} Case Competition','Finalist','',dict(sy=gy-2,sm=2,ey=gy-2,em=2),['Placed 2nd of 40 teams with a supply-chain turnaround plan'])
        if r.random()<0.15:add('lead','Stock pitch competition','Undergraduate Stock Pitch Competition','Semifinalist','',dict(sy=gy-2,sm=11,ey=gy-2,em=11),['Pitched a long thesis on a mid-cap retailer to a panel of 5 investors'])
        if r.random()<0.15:add('lead','Greek life',r.choice(['Kappa Alpha Theta','Sigma Chi','Alpha Kappa Psi','Delta Sigma Pi']),r.choice(['Vice President of Finance','Member','Philanthropy Chair']),'',dict(sy=gy-3,sm=1,cur=True),['Managed a $40K chapter budget'])
        if r.random()<0.1:add('lead','Military / ROTC','Army ROTC','Cadet','',dict(sy=gy-4,sm=8,cur=True),['Completed leadership training with a 30-cadet platoon'])
    if level=='athlete':
        sport,lv=r.choice(SPORTS);add('ath','Varsity athlete',f'{short} Varsity {sport}',r.choice(['Team Captain','Student-Athlete']),'',dict(sy=gy-4,sm=8,cur=True),[f'Competed at the {lv} level while training 20 hours per week','Named to the conference All-Academic Team'])
    if r.random()<0.35:vo,vt=r.choice(VOL);vy=min(gy-4,TY-2);add('vol','Volunteer',vo,vt,f'{sc}, {sst}',dict(sy=vy,sm=1,ey=vy+1,em=12),['Built 3 homes with a 12-person crew' if 'Habitat' in vo else 'Sorted and packed 2,000 lbs of food monthly'])
    if field=='tech':
        add('proj','Project','',r.choice(['Portfolio Tracker','Campus Ride Share App','Course Planner']),'',dict(sy=gy-2,sm=1,ey=gy-2,em=4),['Built a web app with React and Flask used by 150 students'])
        if r.random()<0.3:add('proj','Hackathon','HackGT','Participant','',dict(sy=gy-2,sm=10,ey=gy-2,em=10),['Built a transit alert bot in a 36-hour hackathon; won Best Use of Data'])
        if r.random()<0.3:add('proj','Open source','pandas','Open Source Contributor','',dict(sy=gy-2,sm=6,cur=True),['Merged 4 pull requests fixing datetime parsing bugs'])
    if level=='dense':
        for k in range(2):o,c,s=r.choice(P['orgs']);tt=r.choice(P['titles']);add('exp',ityp(tt),o,tt,f'{c}, {s}',dict(sy=gy-3-k,sm=6,ey=gy-3-k,em=8),r.sample(P['bullets'],3))
    # skills / certs / langs
    p['skills']=r.sample(P['skills'],r.randint(4,min(7,len(P['skills']))))
    if r.random()<0.2:p['skills'].append(('Kanban','Kanban'))
    if r.random()<0.25:p['soft']=r.sample(['Public Speaking','Teamwork','Time Management','Leadership'],2)
    nc={'dense':3,'freshman':0}.get(level,r.choice([0,1,1,2]))
    p['certs']=r.sample(P['certs'],min(nc,len(P['certs'])))
    if r.random()<0.1:p['certs'].append(('Dale Carnegie Leadership Training Certificate','Dale Carnegie Leadership Training Certificate'))
    if level=='international' or r.random()<0.35:p['langs']=r.sample(LANGS,r.choice([1,2]))
    if r.random()<0.5:p['interests']=r.sample(INTERESTS,3)
    if level=='international':p['workauth']=('Work Authorization: Authorized to work in the U.S.; no sponsorship required','Authorized to work in the U.S. (no sponsorship needed)')
    elif r.random()<0.15:p['workauth']=('Work Authorization: U.S. Citizen','U.S. citizen')
    if 'pubs' in edges:p['pubs']=[f'{last}, {first[0]}., & Rivera, M. (2026). Sleep and study habits in first-year students. Undergraduate Research Journal, 4(2), 11–19.']
    # no future dates: anything ending after today is ongoing; anything starting after today is dropped
    p['exp']=[x for x in p['exp'] if (x['d']['sy'],x['d']['sm'])<=(TY,TM)]
    for x in p['exp']:
        d=x['d']
        if not d.get('cur') and (d['ey'],d['em'])>(TY,TM):d['cur']=True;d.pop('season',None)
    # dates printed
    for x in p['exp']:
        pr,tr=fmt_dates(x['d'],x['style']);x['dp']=pr;x['dt']=tr
    return p
def truth(p):
    edu=[]
    for e in p['edu']:
        edu.append(dict(school=e['school'],ambiguous=e['ambiguous'],degree=e['degree'],majors=e['majors'],minors=e['minors'],gpa=e['gpa'] or '',gpaScale=e['scale'] if e['gpa'] else '',majorGpa=e['mgpa'],gradYear=str(e['gy']),gradMonth=e['gm'],classYear=e['classYear']))
    return {'_synthetic':True,'_note':'Fictional test person. Not a real resume.','id':p['id'],'format':p['fmt'],'field':p['field'],'level':p['level'],'edges':p['edges'],'today':'%d-%02d-%02d'%TODAY,
      'contact':dict(name=p['name'],email=p['email'],phone=p['phone'],location=f"{p['city']}, {p['st']}"),'education':edu,
      'experience':[dict(type=x['type'],org=x['org'],title=x['title'],location=x['loc'],start=x['dt']['start'],end=x['dt']['end'],months=x['dt']['months']) for x in p['exp']],
      'certs':[c for _,c in p['certs']],'skills':[c for _,c in p['skills']],'languages':[dict(name=n,level=l) for n,l in p['langs']],'workAuth':p['workauth'][1]}
# ---------------- document model ----------------
SECORD=[('exp','exp'),('proj','proj'),('lead','lead'),('ath','ath'),('vol','vol')]
def model(p,r):
    """Returns ordered list of (kind, payload). kinds: name, contact, head, entry(l1L,l1R,l2L,l2R,bul,extra), line"""
    M=[('name',p['name']),('contact',[f"{p['city']}, {p['st']}",p['email'],p['phone'],p['linkedin']])]
    side=[]
    M.append(('head',r.choice(HEAD['edu']),'edu'))
    for e in p['edu']:
        dl=e['deg_print']+('' if not e['majors'] else (' in ' if e['deg_print'].startswith(('Bachelor','Associate')) else ', ')+' and '.join(e['majors']))
        if e['level']=='hs':dl='High School Diploma'
        extra=[]
        g=[]
        if e['gpa']:g.append(f"GPA: {e['gpa']}"+(f"/{e['scale']}" if e['scale'] else '') if not e['mgpa'] else f"Cumulative GPA: {e['gpa']}"+(f"/{e['scale']}" if e['scale'] else ''))
        if e['mgpa']:g.append(f"Major GPA: {e['mgpa']}")
        if e['minors']:g.append('Minor: '+', '.join(e['minors']))
        if e['honors']:g.append('; '.join(e['honors']))
        if g:extra.append(' | '.join(g))
        if e['course']:extra.append('Relevant Coursework: '+', '.join(e['course']))
        if e['abroad']:extra.append(f"{e['abroad'][0]} – {e['abroad'][1]}")
        M.append(('entry',dict(l1=(e['school_print'],e['loc']),l2=(dl,e['gp']),bul=[],extra=extra,kind='edu')))
    for sk,hk in SECORD:
        xs=[x for x in p['exp'] if x['sec']==sk]
        if not xs:continue
        M.append(('head',r.choice(HEAD[hk]),sk))
        for x in xs:M.append(('entry',dict(l1=(x['org'] if x['org'] else x['title'],x['loc']),l2=(x['title'] if x['org'] else '',x['dp']),bul=x['bullets'],extra=[],kind=sk,x=x)))
    if p['pubs']:M.append(('head','PUBLICATIONS','pub'));[M.append(('line',t)) for t in p['pubs']]
    sk=[('Technical Skills' if p['field'] in('tech','engineering') else r.choice(['Technical','Software','Skills','Computer']),', '.join(s for s,_ in p['skills']))]
    if p['soft']:sk.append(('Soft Skills',', '.join(p['soft'])))
    if p['certs']:sk.append(('Certifications',', '.join(c for c,_ in p['certs'])))
    if p['langs']:sk.append(('Languages',', '.join(f'{n} ({l})' for n,l in p['langs'])))
    if p['interests']:sk.append(('Interests',', '.join(p['interests'])))
    if p['workauth'][0]:sk.append(tuple(p['workauth'][0].split(': ',1)))
    M.append(('head',r.choice(HEAD['skills']),'skills'))
    for a,b in sk:M.append(('line',f'{a}: {b}'))
    return M
# ---------------- PDF renderers ----------------
class Pdf:
    def __init__(s,path,name,fonts=('LS','LSB','LSI'),margin=0.6*72,footer=False,lig=False):
        s.c=canvas.Canvas(path,pagesize=letter);s.c.setTitle('SYNTHETIC TEST RESUME – fictional person');s.W,s.H=letter;s.m=margin;s.y=s.H-margin;s.f=fonts;s.page=1;s.footer=footer;s.name=name;s.lig=lig;s.x0=margin;s.x1=s.W-margin
    def L(s,t):return t.replace('fi','\ufb01').replace('fl','\ufb02') if s.lig else t
    def need(s,h):
        if s.y-h<s.m+(14 if s.footer else 0):s.newpage()
    def newpage(s):
        s.foot();s.c.showPage();s.page+=1;s.y=s.H-s.m
    def foot(s):
        if s.footer:s.c.setFont(s.f[0],8);s.c.drawCentredString(s.W/2,s.m-10,f'{s.name} · Page {s.page} of 2')
    def text(s,t,size=10,font=0,x=None,right=None,cs=0,color=black,adv=True):
        s.need(size+3);s.c.setFillColor(color);fn=s.f[font];x=s.x0 if x is None else x
        if cs:
            o=s.c.beginText(x,s.y-size);o.setFont(fn,size);o.setCharSpace(cs);o.textOut(t);o.setCharSpace(0);s.c.drawText(o)
        else:s.c.setFont(fn,size);s.c.drawString(x,s.y-size,t)
        if right:s.c.setFont(s.f[2] if font==2 else s.f[0],size);s.c.drawRightString(s.x1,s.y-size,right)
        s.c.setFillColor(black)
        if adv:s.y-=size+3
    def wrap(s,t,size,width,font=0):
        words=t.split(' ');lines=[];cur=''
        for w in words:
            tr=(cur+' '+w).strip()
            if pdfmetrics.stringWidth(tr,s.f[font],size)>width and cur:lines.append(cur);cur=w
            else:cur=tr
        if cur:lines.append(cur)
        return lines
    def bullet(s,t,size=9.5,ind=10,glyph='•'):
        ls=s.wrap(s.L(t),size,s.x1-s.x0-ind-8)
        for k,l in enumerate(ls):
            s.need(size+2)
            if k==0:s.c.setFont(s.f[0],size);s.c.drawString(s.x0+ind-8,s.y-size,glyph)
            s.c.setFont(s.f[0],size);s.c.drawString(s.x0+ind+2,s.y-size,l);s.y-=size+2
    def para(s,t,size=9.5,font=0):
        for l in s.wrap(t,size,s.x1-s.x0,font):s.text(l,size,font)
    def rule(s):s.c.setLineWidth(0.6);s.c.line(s.x0,s.y+1,s.x1,s.y+1);s.y-=4
    def save(s):s.foot();s.c.save()
def r_onecol(p,M,path,jake=False):
    two='twopage' in p['edges']
    d=Pdf(path,p['name'],fonts=('SR','SRB','SRI') if jake else ('LS','LSB','LSI'),footer=two,lig=jake)
    for k in M:
        if k[0]=='name':d.c.setFont(d.f[1],22 if jake else 18);d.c.drawCentredString(d.W/2,d.y-20,k[1].upper() if not jake else k[1]);d.y-=28
        elif k[0]=='contact':d.c.setFont(d.f[0],9);d.c.drawCentredString(d.W/2,d.y-9,' | '.join(k[1]) if jake else '  •  '.join(k[1]));d.y-=18
        elif k[0]=='head':
            if 'noheadings' in p['edges']:d.y-=10;continue
            d.y-=6;d.need(30);d.text(k[1].upper() if jake else k[1].upper(),11.5 if jake else 11,1);d.rule()
        elif k[0]=='entry':
            e=k[1];d.need(40)
            if jake and e['kind']!='edu':
                a,b=e['l1'],e['l2']
                if b[0]:d.text(b[0],10,1,right=b[1]);d.text(a[0],9.5,2,right=a[1] or None)
                else:d.text(a[0],10,1,right=b[1])
            else:
                d.text(e['l1'][0],10,1,right=e['l1'][1] or None)
                if e['l2'][0]:d.text(e['l2'][0],9.5,2,right=e['l2'][1] or None)
                elif e['l2'][1]:d.text('',9.5,2,right=e['l2'][1])
            for x in e['extra']:d.para(x,9.5)
            for b in e['bul']:d.bullet(b)
            d.y-=3
        elif k[0]=='line':d.para(d.L(k[1]),9.5)
    d.save()
def r_table(p,M,path):
    d=Pdf(path,p['name'],fonts=('LS','LSB','LSI'));dc=d.x0+90
    for k in M:
        if k[0]=='name':d.text(k[1],20,1);d.y-=2
        elif k[0]=='contact':d.text(' | '.join(k[1]),9);d.y-=8
        elif k[0]=='head':d.y-=4;d.need(30);d.text(k[1].upper(),11,1);d.rule()
        elif k[0]=='entry':
            e=k[1];d.need(40);y0=d.y;dates=e['l2'][1]
            d.c.setFont('LS',8.5);d.c.drawString(d.x0,d.y-10,dates)
            ox=d.x0;d.x0=dc
            d.text(e['l1'][0]+(', '+e['l1'][1] if e['l1'][1] else ''),10,1)
            if e['l2'][0]:d.text(e['l2'][0],9.5,2)
            for x in e['extra']:d.para(x,9.5)
            for b in e['bul']:d.bullet(b)
            d.x0=ox;d.y-=3
        elif k[0]=='line':d.para(k[1],9.5)
    d.save()
def split_side(M):
    main=[];side=[];sec=None
    for k in M:
        if k[0]=='head':sec=k[2]
        if k[0] in('name',):main.append(k)
        elif k[0]=='contact' or sec=='skills':side.append(k)
        else:main.append(k)
    return main,side
def r_twocol(p,M,path,canva=False):
    d=Pdf(path,p['name'],fonts=('LS','LSB','LSI'));W=d.W;main,side=split_side(M)
    if canva:
        d.c.setFillColor(HexColor('#1f3a5f'));d.c.rect(0,d.H-95,W,95,stroke=0,fill=1)
        d.c.setFillColor(white);d.c.setFont('LSB',26);d.c.drawString(d.m,d.H-50,p['name'].upper());d.c.setFont('LS',11);d.c.drawString(d.m,d.H-72,{'finance':'Finance Student','tech':'Computer Science Student','consulting':'Aspiring Consultant','marketing':'Marketing Student','engineering':'Engineering Student','health':'Health Sciences Student','libarts':'Liberal Arts Student'}[p['field']]);d.c.setFillColor(black)
        d.y=d.H-115;sx0,sx1,mx0,mx1=W-d.m-170,W-d.m,d.m,W-d.m-195
    else:
        d.c.setFillColor(HexColor('#eef1f5'));d.c.rect(0,0,205,d.H,stroke=0,fill=1);d.c.setFillColor(black)
        sx0,sx1,mx0,mx1=24,185,230,W-d.m
    ytop=d.y
    # sidebar
    d.x0,d.x1=sx0,sx1;d.y=ytop
    if not canva:d.text(p['name'].split()[0].upper(),20,1);d.text(p['name'].split()[1].upper(),20,1);d.y-=8
    for k in side:
        if k[0]=='contact':
            d.text('C O N T A C T' if canva else 'CONTACT',10.5,1,cs=0);[d.text(c,8.5) for c in k[1]];d.y-=6
        elif k[0]=='head':d.text('S K I L L S' if canva else 'SKILLS',10.5,1);d.y-=1
        elif k[0]=='line':
            a,b=k[1].split(': ',1)
            if a in('Certifications','Languages','Interests','Work Authorization'):
                d.y-=5;d.text(a.upper(),9.5,1)
                items=b.split(', ') if a!='Work Authorization' else [b]
                if a=='Languages':items=[x.replace(' (',' – ').rstrip(')') for x in items]
                if a=='Work Authorization':items=['Work Authorization: '+b]
                for it in items:
                    for l in d.wrap(it,8.5,sx1-sx0):d.text(l,8.5)
            else:
                d.y-=3;d.text(a+':',9,1)
                for it in b.split(', '):d.text(it,8.5)
    # main
    d.x0,d.x1=mx0,mx1;d.y=ytop
    for k in main:
        if k[0]=='name':continue
        if k[0]=='head':
            d.y-=6;t=k[1].upper()
            if canva:d.text(t,11,1,cs=2.2,color=HexColor('#1f3a5f'))
            else:d.text(t,11,1)
            d.rule()
        elif k[0]=='entry':
            e=k[1];d.text(e['l1'][0],10,1)
            sub=' | '.join(x for x in [e['l2'][0],e['l1'][1],e['l2'][1]] if x)
            d.text(sub,8.8,2)
            for x in e['extra']:d.para(x,8.8)
            for b in e['bul']:d.bullet(b,8.8)
            d.y-=3
        elif k[0]=='line':d.para(k[1],8.8)
    d.save()
def text_doc(p,M,r):
    out=[];inline=r.random()<0.5
    for k in M:
        if k[0]=='name':out.append(k[1].upper())
        elif k[0]=='contact':out.append(' | '.join(k[1]));out.append('')
        elif k[0]=='head':
            out.append('')
            if 'noheadings' not in p['edges']:out.append(k[1].upper())
        elif k[0]=='entry':
            e=k[1]
            if inline and e['kind']!='edu':out.append(' | '.join(x for x in [e['l2'][0],e['l1'][0],e['l1'][1],e['l2'][1]] if x))
            else:
                out.append(e['l1'][0]+('      '+e['l1'][1] if e['l1'][1] else ''))
                if e['l2'][0] or e['l2'][1]:out.append((e['l2'][0]+'      '+e['l2'][1]).strip())
            out+=e['extra'];out+=['- '+b for b in e['bul']]
            if 'noheadings' in p['edges']:out.append('')
        elif k[0]=='line':out.append(k[1])
    return '\n'.join(out).replace('\n\n\n','\n\n')+'\n'
def r_textpdf(txt,path,name):
    c=canvas.Canvas(path,pagesize=letter);c.setTitle('SYNTHETIC TEST RESUME – fictional person');y=letter[1]-50
    for l in txt.split('\n'):
        import textwrap
        for w in (textwrap.wrap(l.replace('\t','    '),105,subsequent_indent='  ') or ['']):
            if y<50:c.showPage();y=letter[1]-50
            c.setFont('MO',8.6);c.drawString(40,y,w);y-=11
    c.save()
# ---------------- DOCX ----------------
def para(doc,t,bold=False,italic=False,size=10,style=None,tab=None):
    pr=doc.add_paragraph(style=style);pf=pr.paragraph_format;pf.space_after=Pt(0);pf.space_before=Pt(0)
    if tab:pf.tab_stops.add_tab_stop(Inches(7.3),WD_TAB_ALIGNMENT.RIGHT);t=t+'\t'+tab
    run=pr.add_run(t);run.bold=bold;run.italic=italic;run.font.size=Pt(size);return pr
def r_docx(p,M,path,mode,r):
    doc=docx.Document();sec=doc.sections[0];sec.left_margin=sec.right_margin=Inches(0.6);sec.top_margin=sec.bottom_margin=Inches(0.5)
    st=doc.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(10)
    doc.core_properties.title='SYNTHETIC TEST RESUME – fictional person'
    def entry(target,e,tabbed=True):
        if tabbed:
            para(target,e['l1'][0],bold=True,tab=e['l1'][1] or None)
            if e['l2'][0] or e['l2'][1]:para(target,e['l2'][0],italic=True,tab=e['l2'][1] or None)
        else:
            para(target,e['l1'][0],bold=True)
            para(target,' | '.join(x for x in [e['l2'][0],e['l1'][1],e['l2'][1]] if x),italic=True,size=9)
        for x in e['extra']:para(target,x)
        for b in e['bul']:para(target,b,style='List Bullet')
    if mode in('twocol','canva'):
        main,side=split_side(M)
        para(doc,p['name'].upper(),bold=True,size=20)
        t=doc.add_table(rows=1,cols=2);a,b=t.rows[0].cells
        if mode=='canva':a,b=b,a
        a.paragraphs[0].text='';b.paragraphs[0].text=''
        for k in side:
            if k[0]=='contact':para(a,'CONTACT',bold=True);[para(a,c,size=9) for c in k[1]]
            elif k[0]=='head':para(a,'SKILLS',bold=True)
            elif k[0]=='line':para(a,k[1],size=9)
        for k in main:
            if k[0]=='head':para(b,k[1].upper(),bold=True,size=11)
            elif k[0]=='entry':entry(b,k[1],tabbed=False)
            elif k[0]=='line':para(b,k[1])
    elif mode=='table':
        para(doc,p['name'],bold=True,size=20)
        for k in M:
            if k[0]=='contact':para(doc,' | '.join(k[1]),size=9)
            elif k[0]=='head':para(doc,k[1].upper(),bold=True,size=11)
            elif k[0]=='entry':
                e=k[1];t=doc.add_table(rows=1,cols=2);a,b=t.rows[0].cells;a.width=Inches(1.3);b.width=Inches(6)
                a.paragraphs[0].text=e['l2'][1];b.paragraphs[0].text=''
                para(b,e['l1'][0]+(', '+e['l1'][1] if e['l1'][1] else ''),bold=True)
                if e['l2'][0]:para(b,e['l2'][0],italic=True)
                for x in e['extra']:para(b,x)
                for bb in e['bul']:para(b,bb,style='List Bullet')
            elif k[0]=='line':para(doc,k[1])
    elif mode=='text':
        for l in text_doc(p,M,random.Random(hash(p['id'])%999)).split('\n'):
            if l.startswith('- '):para(doc,l[2:],style='List Bullet')
            else:para(doc,l.replace('      ','\t'))
    else:
        jake=mode=='jake'
        for k in M:
            if k[0]=='name':pr=para(doc,k[1] if jake else k[1].upper(),bold=True,size=20)
            elif k[0]=='contact':para(doc,' | '.join(k[1]),size=9)
            elif k[0]=='head':
                if 'noheadings' in p['edges']:para(doc,'');continue
                para(doc,k[1].upper(),bold=True,size=11.5)
            elif k[0]=='entry':
                e=k[1]
                if jake and e['kind']!='edu' and e['l2'][0]:
                    para(doc,e['l2'][0],bold=True,tab=e['l2'][1]);para(doc,e['l1'][0],italic=True,tab=e['l1'][1] or None)
                    for b in e['bul']:para(doc,b,style='List Bullet')
                else:entry(doc,e)
            elif k[0]=='line':para(doc,k[1])
    doc.save(path)
# ---------------- schedule ----------------
FMTS=['onecol','twocol','table','canva','jake','word','text']
FIELDS=['finance','tech','consulting','marketing','engineering','health','libarts']
LEVELS=(['standard']*24+['freshman']*5+['transfer']*5+['athlete']*6+['international']*5+['career']*5+['dense']*5+['standard']*5)
def edges_for(i,fmt,level):
    r=random.Random(77+i);e=[]
    if i%6==1:e.append('nogpa')
    if i%6==3:e.append('weirddates')
    if i%9==4:e.append('majorgpa')
    if i%11==5:e.append('double')
    if i%13==7:e.append('abroad')
    if fmt in('text','onecol','word') and i%10==6:e.append('noheadings')
    if level=='dense' and fmt in('onecol','jake','word'):e.append('twopage')
    if i in(17,44):e.append('ambiguous')
    if level=='international' and i%2==0:e.append('scale10')
    if i%15==2:e.append('pubs')
    return e
def main():
    import glob
    for d in glob.glob(os.path.join(OUT,'f[0-9][0-9]')):shutil.rmtree(d)
    rows=[]
    for i in range(60):
        fmt=FMTS[i%7];field=FIELDS[(i//7+i)%7];level=LEVELS[(i*37)%60]
        if level=='dense' and fmt not in('onecol','jake','word'):fmt=['onecol','jake','word'][i%3]
        ed=edges_for(i,fmt,level);p=person(i,fmt,field,level,ed)
        if 'twopage' in ed:  # pad dense resumes so they really run to two pages
            P=F[field];r=random.Random(i)
            for o,c,s in r.sample(P['orgs'],3):
                tt=r.choice(P['titles']);x=dict(sec='exp',type=ityp(tt),org=o,title=tt,loc=f'{c}, {s}',d=dict(sy=2022,sm=6,ey=2022,em=8),style='mon',bullets=P['bullets'][:4])
                x['dp'],x['dt']=fmt_dates(x['d'],'mon');p['exp'].append(x)
        d=os.path.join(OUT,p['id']);os.makedirs(d,exist_ok=True)
        r=random.Random(500+i);M=model(p,r)
        pdf=os.path.join(d,'resume.pdf');dx=os.path.join(d,'resume.docx')
        if fmt in('onecol','jake'):r_onecol(p,M,pdf,jake=fmt=='jake')
        elif fmt=='table':r_table(p,M,pdf)
        elif fmt in('twocol','canva'):r_twocol(p,M,pdf,canva=fmt=='canva')
        elif fmt=='text':
            t=text_doc(p,M,random.Random(hash(p['id'])%999));open(os.path.join(d,'resume.txt'),'w').write(t);r_textpdf(t,pdf,p['name'])
        r_docx(p,M,dx,fmt,r)
        if fmt=='word':
            subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',d,dx],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=120)
        T=truth(p);T['files']=sorted(x for x in os.listdir(d) if x.startswith('resume.'))
        json.dump(T,open(os.path.join(d,'expected.json'),'w'),indent=1,ensure_ascii=False)
        rows.append((p['id'],fmt,field,level,','.join(ed)))
    with open(os.path.join(OUT,'INDEX.md'),'w') as f:
        f.write('# Synthetic resume fixtures (V14.3)\n\nAll people are fictional. Generated by `build.py` (seeded). Each folder has `resume.pdf`, `resume.docx` (and `resume.txt` for the text format) plus a hand-specified `expected.json`.\n\n| id | format | field | level | edge cases |\n|---|---|---|---|---|\n')
        for r_ in rows:f.write('| '+' | '.join(r_)+' |\n')
    print(len(rows),'fixtures')
if __name__=='__main__':main()
